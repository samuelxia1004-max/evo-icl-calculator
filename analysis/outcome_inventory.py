"""Aggregate availability of functional/safety outcomes for the next study.

No assumption that the unqualified visual-acuity column means UDVA or CDVA.
No patient identifiers, dates, cell excerpts or individual rows are exported.
"""
import argparse
from pathlib import Path
import json
import numpy as np
from openpyxl import load_workbook
from read_registry import number, numeric_endpoint

parser=argparse.ArgumentParser()
parser.add_argument("--workbook",required=True)
parser.add_argument("--output",default="research/outcome-inventory.json")
args=parser.parse_args()
workbook=load_workbook(args.workbook,read_only=True,data_only=True)
values={key:[] for key in ["va_month1","va_month3","va_month6","va_year1","ecd_month1","ecd_month3","ecd_month6","ecd_year1"]}
total=0
for sheet in workbook:
    if sheet.title not in ("2022","2023","2024","2025"):continue
    header=3 if sheet.title=="2023" else 2
    shift=0 if sheet.title in ("2022","2023") else -2
    for row in sheet.iter_rows(min_row=header+2,values_only=True):
        if not row[1]:continue
        for eye in (0,1):
            if not numeric_endpoint(row[64+shift+eye]) and not numeric_endpoint(row[84+shift+eye]):continue
            total+=1
            columns={"va_month1":77,"va_month3":89,"va_month6":101,"va_year1":113,
                     "ecd_month1":87,"ecd_month3":99,"ecd_month6":111,"ecd_year1":123}
            for key,column in columns.items():values[key].append(number(row[column-1+shift+eye]))
output={"eligible_eyes":total,"variables":{},"visual_acuity_definition":"Unqualified registry visual-acuity columns; UDVA/CDVA and scale require clinician confirmation.",
        "scope":"Feasibility inventory only, not an analysis of functional rehabilitation efficacy."}
for key,array in values.items():
    a=np.asarray(array);valid=a[np.isfinite(a)]
    output["variables"][key]={"numeric_available":len(valid),"missing_or_nonnumeric":int((~np.isfinite(a)).sum()),
        "numeric_min":float(valid.min()) if len(valid) else None,"numeric_max":float(valid.max()) if len(valid) else None}
Path(args.output).parent.mkdir(parents=True,exist_ok=True)
Path(args.output).write_text(json.dumps(output,ensure_ascii=False,indent=2)+"\n")
print("Functional/safety availability inventory saved; no row data exported.")
